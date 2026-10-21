"""A re-screened Capital IQ export that drops or moves columns stops with a
readable message (which columns, which letters), not an IndexError."""
import unittest

import pandas as pd

from nexus.capitaliq_layout import check_or_exit, fill_warnings, problems, resolve, where
from nexus.settings import Settings

GOOD = ["Company Name", "Exchange:Ticker", "Industry Classifications", "Company Type",
        "Geographic Locations", "Total Employees [Latest Annual]",
        "Total Revenue [LTM] ($USDmm, Historical rate)", "Competitors", "Offices",
        "Geographic Region", "Business Description", "Website", "Company Status"]
# the October 2026 re-screen: no description, other order, 10 columns
NEW = ["Company Name", "Exchange:Ticker", "Industry Classifications", "Company Type",
       "Geographic Locations", "Company Status", "Revenues [LTM] ($USDmm, Historical rate)",
       "Website", "Number of Employees [Latest Annual]", "Offices"]


class TestLayout(unittest.TestCase):

    def setUp(self):
        self.ciq = Settings().capitaliq

    def test_good_export(self):
        self.assertEqual(problems(GOOD, self.ciq), [])

    def test_rescreened_export(self):
        p = " ".join(problems(NEW, self.ciq))
        self.assertIn("col_description: column K does not exist", p)
        self.assertIn("col_status: column M does not exist", p)
        w = where(NEW)
        self.assertIsNone(w["col_description"])
        self.assertEqual((w["col_status"], w["col_website"], w["col_address"]), ("F", "H", "J"))
        with self.assertRaises(SystemExit) as e:
            check_or_exit(NEW, self.ciq, "companies_raw.xlsx")
        msg = str(e.exception)
        self.assertIn("Re-export with these data items: Business Description", msg)
        self.assertIn("col_status: F", msg)

    def test_empty_revenue_is_flagged(self):
        df = pd.DataFrame([["A", "-", "x", "y", "z", "5", "-", "", "addr", "", "desc", "w", "Operating"]] * 20,
                          columns=GOOD)
        w = " ".join(fill_warnings(df, self.ciq))
        self.assertIn("col_revenue", w)
        self.assertNotIn("col_description", w)


# the Pharma export: an extra revenue item, everything after it shifted
SHIFTED = ["Company Name", "Exchange:Ticker", "Industry Classifications", "Company Type",
           "Geographic Locations", "Total Employees [Latest Annual]",
           "Total Revenue [LTM] ($USDmm, Historical rate)",
           "Total Revenue [Latest Annual] ($USDmm, Historical rate)", "Competitors", "Offices",
           "Geographic Region", "Business Description", "Website", "Company Status"]


class TestFoundByName(unittest.TestCase):

    def setUp(self):
        self.ciq = Settings().raw["capitaliq"]

    def test_shifted_columns_are_found(self):
        used, moved = resolve(SHIFTED, self.ciq)
        self.assertEqual(problems(SHIFTED, used), [])
        self.assertEqual((used["col_revenue"], used["col_address"], used["col_description"],
                          used["col_website"], used["col_status"]), ("G", "J", "L", "M", "N"))
        self.assertIn("col_status", moved)

    def test_ltm_revenue_preferred(self):
        h = list(SHIFTED)
        h[6], h[7] = h[7], h[6]                 # Latest Annual first, LTM second
        self.assertEqual(resolve(h, self.ciq)[0]["col_revenue"], "H")

    def test_any_order(self):
        h = list(reversed(GOOD))
        used, _ = resolve(h, self.ciq)
        self.assertEqual(problems(h, used), [])
        self.assertEqual(used["col_company"], "M")

    def test_missing_column_still_reported(self):
        used, _ = resolve(NEW, self.ciq)
        self.assertIn("col_description", " ".join(problems(NEW, used)))

    def test_settings_uses_the_export_header(self):
        import tempfile, yaml
        from pathlib import Path
        from openpyxl import Workbook
        d = Path(tempfile.mkdtemp())
        wb = Workbook(); ws = wb.active
        for j, h in enumerate(SHIFTED, start=1):
            ws.cell(8, j, h)
        ws.cell(9, 1, "Acme Inc.")
        wb.save(d / "raw.xlsx")
        cfg = Settings().raw
        cfg["paths"] = dict(cfg["paths"], companies_raw=str(d / "raw.xlsx"))
        (d / "config.yaml").write_text(yaml.safe_dump(cfg))
        ciq = Settings(d / "config.yaml").capitaliq
        self.assertEqual((ciq["col_description"], ciq["col_status"]), ("L", "N"))


if __name__ == "__main__":
    unittest.main()
