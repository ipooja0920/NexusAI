"""Faculty publications: titles weighted by age and type, a pre-filter
search of their own, and the top titles in the profile the AI reads."""
import csv
import tempfile
import unittest
from pathlib import Path

import openpyxl

from nexus.prefilter import rank_by_weighted_queries
from nexus.publications import (age_weight, find_publications, profile_block,
                                read_publications, settings, type_weight)
from tests.harness import TempProject, install_fake_openai

HEADER = ["NetID", "Name", "Department", "Title", "Type", "Status", "Year", "Description"]


def write(path, rows):
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(HEADER)
        for r in rows:
            w.writerow(r)


def row(name, title, kind="Journal Article", year="2020"):
    return ["x1", name, "Dept", title, kind, "", year, "free text never read"]


class TestWeights(unittest.TestCase):

    def setUp(self):
        self.s = settings({})

    def test_age(self):
        self.assertEqual([age_weight(y, self.s, today=2026) for y in (2025, 2021, 2016, 2011, 2001)],
                         [1.0, 1.0, 0.6, 0.4, 0.25])
        self.assertEqual(age_weight(None, self.s, today=2026), 1.0)

    def test_type(self):
        self.assertEqual(type_weight("Journal Article", "X", self.s), 1.0)
        self.assertEqual(type_weight("Review Article", "X", self.s), 1.0)
        self.assertEqual(type_weight("Poster/Presentation", "X", self.s), 0.5)
        self.assertEqual(type_weight("Journal Erratum", "X", self.s), 0.0)
        self.assertEqual(type_weight("Journal Article", "Erratum: something", self.s), 0.0)

    def test_config_overrides(self):
        s = settings({"publications": {"older": 0.1}})
        self.assertEqual(age_weight(1990, s, today=2026), 0.1)


class TestReading(unittest.TestCase):

    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())
        self.s = settings({})

    def test_errata_duplicates_and_other_people(self):
        write(self.dir / "a.csv", [
            row("Leighton Core", "PRO-seq maps RNA polymerase", year="2016"),
            row("Leighton Core", "PRO-seq maps RNA polymerase", year=""),       # stray copy
            row("Leighton Core", "Erratum to PRO-seq", kind="Journal Erratum", year="2017"),
            row("Leighton Core", "Centromere transcription", year="2024"),
            row("Leighton Core", "Nascent RNA poster", kind="Poster/Presentation", year="2024"),
            row("Someone Else", "Soil science", year="2025"),
        ])
        p = read_publications(self.dir / "a.csv", "leighton  core", self.s, today=2026)
        self.assertEqual([x["title"] for x in p],
                         ["Centromere transcription", "PRO-seq maps RNA polymerase",
                          "Nascent RNA poster"])
        self.assertEqual([x["weight"] for x in p], [1.0, 0.6, 0.5])
        self.assertEqual(p[1]["year"], 2016)          # the dated copy is kept
        self.assertNotIn("free text", str(p))

    def test_find_across_files_and_missing_folder(self):
        write(self.dir / "a.csv", [row("Ann Bee", "Topic A", year="2025")])
        write(self.dir / "b.csv", [row("Ann Bee", "Topic B", year="2010")])
        (self.dir / "notes.csv").write_text("just,some\n1,2\n")       # not a publications file
        self.assertEqual([p["title"] for p in find_publications(self.dir, "Ann Bee", self.s, 2026)],
                         ["Topic A", "Topic B"])
        self.assertEqual(find_publications(self.dir / "nope", "Ann Bee", self.s), [])

    def test_profile_block(self):
        pubs = [{"title": "New", "year": None, "weight": 1.0},
                {"title": "Old", "year": 2001, "weight": 0.25}]
        b = profile_block(pubs, 1)
        self.assertIn("the keywords and description above come first", b)
        self.assertIn("- forthcoming: New", b)
        self.assertNotIn("Old", b)
        self.assertEqual(profile_block([], 12), "")


class TestWeightedSearch(unittest.TestCase):

    def test_best_weighted_match(self):
        companies = [[1, 0], [0, 1], None]
        # title 1 points at company 0 but is old (0.25); title 2 at company 1, recent
        s = rank_by_weighted_queries([[1, 0], [0, 1]], [0.25, 1.0], companies)
        self.assertGreater(s[1], s[0])
        self.assertAlmostEqual(s[0], 0.25, places=5)
        self.assertEqual(s[2], -1.0)
        self.assertEqual(rank_by_weighted_queries([], [], companies), [-1.0] * 3)


class TestInARun(unittest.TestCase):

    def test_publications_reach_the_profile_and_run_info(self):
        install_fake_openai()
        p = TempProject(config_overrides={"prefilter": {"top_n_candidates": 3,
                                                        "extra_per_search": 2}})
        p.__enter__()
        try:
            folder = p.dir / "data" / "input" / "publications"
            folder.mkdir(parents=True, exist_ok=True)
            write(folder / "valla.csv", [
                row("Ioulia Valla", "Zeolite catalysts for biomass pyrolysis", year="2024"),
                row("Ioulia Valla", "An erratum", kind="Journal Erratum", year="2024"),
            ])
            p.seed_geocodes()
            p.run_enrich()
            self.assertEqual(p.run_match(["--faculty", "Ioulia Valla"]), 0)
            wb = openpyxl.load_workbook(p.outputs()[0])
            info = {r[0]: r[1] for r in wb["Run Info"].iter_rows(values_only=True)}
            self.assertIn("- 2024: Zeolite catalysts for biomass pyrolysis",
                          info["Research profile used"])
            self.assertNotIn("erratum", info["Research profile used"].lower())
            self.assertTrue(str(info["Publications used"]).startswith("1 title ("))
            self.assertIn("publications search added", info["Publications used"])
        finally:
            p.__exit__()

    def test_no_publications(self):
        install_fake_openai()
        p = TempProject()
        p.__enter__()
        try:
            p.seed_geocodes()
            p.run_enrich()
            self.assertEqual(p.run_match(["--faculty", "Ioulia Valla"]), 0)
            info = {r[0]: r[1] for r in openpyxl.load_workbook(p.outputs()[0])["Run Info"]
                    .iter_rows(values_only=True)}
            self.assertEqual(info["Publications used"], "none found")
            self.assertNotIn("Selected publications", info["Research profile used"])
        finally:
            p.__exit__()


if __name__ == "__main__":
    unittest.main()


class TestAccentedNames(unittest.TestCase):
    """The registration form says 'Cristina Colón-Semenza', her publications
    export 'Cristina Colon-Semenza' - the same person."""

    def test_name_key(self):
        from nexus.faculty_registration import name_key
        self.assertEqual(name_key("Cristina Colón-Semenza"), name_key("cristina  colon semenza"))
        self.assertEqual(name_key("Cristina Colón-Semenza"), "cristinacolonsemenza")
        self.assertNotEqual(name_key("Na Li"), name_key("Na Zhang"))

    def test_publications_found_without_the_accent(self):
        d = Path(tempfile.mkdtemp())
        write(d / "c.csv", [row("Cristina Colon-Semenza", "Telerehabilitation in Parkinson disease",
                                year="2026")])
        p = find_publications(d, "Cristina Colón-Semenza", settings({}), today=2026)
        self.assertEqual([x["title"] for x in p], ["Telerehabilitation in Parkinson disease"])

    def test_faculty_flag_without_the_accent(self):
        import argparse
        from match_faculty import pick_faculty
        install_fake_openai()
        p = TempProject()
        p.__enter__()
        try:
            p.write_registration([{"name": "Cristina Colón-Semenza",
                                   "keywords": "Parkinson disease, physical activity"}])
            cfg = p.settings()
            args = argparse.Namespace(row=None, faculty="Cristina Colon-Semenza")
            todo, err = pick_faculty(args, cfg)
            self.assertIsNone(err)
            self.assertEqual(todo[0]["name"], "Cristina Colón-Semenza")
        finally:
            p.__exit__()
