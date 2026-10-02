"""The BD-IFR Faculty Registration CSV as a faculty source: research fields
read (never emails), a labelled research profile, sector match against
Capital IQ's industry groups, and a full match run by --row."""
import unittest

import openpyxl

from nexus.faculty_registration import read_registration, research_profile, sector_match
from tests.harness import FakeOpenAI, TempProject, install_fake_openai

JULIE = {"name": "Julie Burland", "email": "someone@example.edu", "school": "CAHNR",
         "department": "Kinesiology", "nature": "Applied / Translational, Clinical",
         "keywords": "biomechanics, wearables, zeolite catalysis",
         "research": "Recovery after ACL reconstruction; pyrolysis sensors.",
         "sector1": "Health Care: Equipment and Services",
         "sector2": "Health Care: Pharmaceuticals, Biotechnology and Life Sciences",
         "notes": "Interested in companies testing wearables in our lab."}
OTHER = dict(JULIE, name="Someone Else", keywords="soil", research="Soil science")


class TestReader(unittest.TestCase):

    def test_rows_numbered_like_a_spreadsheet_and_no_email(self):
        with TempProject() as p:
            p.write_registration([OTHER, JULIE])
            fac = read_registration(p.settings().path("faculty_registration"))
            self.assertEqual([(f["row"], f["name"]) for f in fac],
                             [(2, "Someone Else"), (3, "Julie Burland")])
            self.assertNotIn("email", fac[1])
            self.assertNotIn("example.edu", str(fac[1]))

    def test_profile_is_labelled(self):
        with TempProject() as p:
            p.write_registration([JULIE])
            prof = research_profile(read_registration(p.settings().path("faculty_registration"))[0])
        self.assertIn("Research keywords: biomechanics", prof)
        self.assertIn("Nature of research: Applied / Translational, Clinical", prof)
        self.assertIn("Primary industry sector: Health Care: Equipment and Services", prof)
        self.assertIn("Industry interests and goals: Interested in companies testing", prof)
        self.assertNotIn("example.edu", prof)


class TestSectorMatch(unittest.TestCase):

    def test_capital_iq_wording(self):
        self.assertEqual(sector_match("Health Care Equipment and Services (Primary)", JULIE),
                         "Primary sector")
        self.assertEqual(sector_match("Pharmaceuticals, Biotechnology and Life Sciences (Primary)",
                                      JULIE), "Secondary sector")

    def test_other_sectors_do_not_match(self):
        energy = {"sector1": "Energy: Energy Equipment and Services"}
        self.assertEqual(sector_match("Health Care Equipment and Services (Primary)", energy), "")
        self.assertEqual(sector_match(None, JULIE), "")


class TestRunByRow(unittest.TestCase):

    def setUp(self):
        install_fake_openai()

    def test_row_run_two_tabs_with_justification(self):
        p = TempProject(config_overrides={"prefilter": {"top_n_candidates": 3, "extra_per_search": 0}})
        p.__enter__()
        try:
            p.seed_geocodes()
            p.write_registration([OTHER, JULIE])
            p.run_enrich()
            FakeOpenAI.reset()
            flags_before = p.faculty_flags()
            self.assertEqual(p.run_match(["--row", "3"]), 0)
            out = p.outputs()
            self.assertEqual(len(out), 1)
            self.assertIn("Julie Burland", out[0].name)
            wb = openpyxl.load_workbook(out[0])
            self.assertEqual(wb.sheetnames[:3], ["Small - 3 pre-filter",
                                                  "Small - all companies", "Mid ($100m-$1bn)"])
            for name in wb.sheetnames[:3]:
                ws = wb[name]
                hdr = [c.value for c in ws[1]]
                j = hdr.index("Justification") + 1
                self.assertTrue(all(ws.cell(r, j).value for r in range(2, ws.max_row + 1)))
            self.assertGreater(FakeOpenAI.counts().get("justification", 0), 0)
            # the registration run never touches the Faculty Database flags
            self.assertEqual(p.faculty_flags(), flags_before)
            # and no email anywhere in the workbook
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    self.assertFalse(any("example.edu" in str(v) for v in row if v))
        finally:
            p.__exit__()

    def test_faculty_by_name_finds_the_registration(self):
        with TempProject() as p:
            p.seed_geocodes()
            p.write_registration([JULIE])
            p.run_enrich()
            self.assertEqual(p.run_match(["--faculty", "julie  burland"]), 0)

    def test_unknown_row_errors_cleanly(self):
        with TempProject() as p:
            p.seed_geocodes()
            p.write_registration([JULIE])
            p.run_enrich()
            self.assertEqual(p.run_match(["--row", "40"]), 1)


if __name__ == "__main__":
    unittest.main()
