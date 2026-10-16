"""Alignment asked several times per company, middle score used: one ask of
gpt-5-nano can swing by several points (8 one run, 4 the next)."""
import json
import unittest

import openpyxl

from match_faculty import combine_votes
from nexus.scoring import STD_NA_REASON
from tests.harness import FakeOpenAI, TempProject, install_fake_openai


def unreadable(path) -> int:
    """Companies whose answers could not be read - never cached, so every
    run asks again (the fixture has one: Broken)."""
    ws = openpyxl.load_workbook(path)["All Companies"]
    hdr = [c.value for c in ws[1]]
    rc = hdr.index("Alignment Reason")
    return sum(1 for r in ws.iter_rows(min_row=2, values_only=True) if r[rc] == STD_NA_REASON)


def v(score, reason=None):
    return {"score": str(score), "reason": reason or f"r{score}"}


class TestCombineVotes(unittest.TestCase):

    def test_middle_score_outvotes_an_odd_one(self):
        r = combine_votes([v(8), v(4), v(7)])
        self.assertEqual(r["score"], "7")
        self.assertEqual(r["reason"], "r7")
        self.assertEqual(r["votes"], "8, 4, 7")

    def test_even_count_takes_lower_middle(self):
        self.assertEqual(combine_votes([v(8), v(5)])["score"], "5")

    def test_na(self):
        self.assertEqual(combine_votes([v("NA"), v("NA"), v(6)])["score"], "NA")
        self.assertEqual(combine_votes([v("NA"), v(6), v(8)])["score"], "6")
        self.assertEqual(combine_votes([v(9)])["score"], "9")


class TestVotesInARun(unittest.TestCase):

    def test_three_asks_and_votes_column(self):
        install_fake_openai()
        p = TempProject(config_overrides={"openai": {"alignment_votes": 3}})
        p.__enter__()
        try:
            p.seed_geocodes()
            p.run_enrich()
            FakeOpenAI.reset()
            self.assertEqual(p.run_match(["--faculty", "Ioulia Valla"]), 0)
            three = FakeOpenAI.counts().get("alignment", 0)
            wb = openpyxl.load_workbook(p.outputs()[0])
            ws = wb["All Companies"]
            hdr = [c.value for c in ws[1]]
            self.assertIn("Alignment Votes", hdr)
            vc, sc = hdr.index("Alignment Votes"), hdr.index("Alignment Score (1-9)")
            scored = [r for r in ws.iter_rows(min_row=2, values_only=True)
                      if r[sc] not in (None, "NA")]
            self.assertTrue(scored)
            for r in scored:
                votes = str(r[vc]).split(", ")
                self.assertEqual(len(votes), 3)
                nums = sorted(int(x) for x in votes if x != "NA")
                self.assertEqual(int(r[sc]), nums[(len(nums) - 1) // 2])
            self.assertEqual(three % 3, 0)
            # a re-run is all cache hits, except the unreadable answers
            broken = unreadable(p.outputs()[0])
            FakeOpenAI.reset()
            self.assertEqual(p.run_match(["--faculty", "Ioulia Valla"]), 0)
            self.assertEqual(FakeOpenAI.counts().get("alignment", 0), 3 * broken)
            info = {r[0]: r[1] for r in wb["Run Info"].iter_rows(values_only=True)}
            self.assertTrue(str(info["Alignment asks per company"]).startswith("3"))
        finally:
            p.__exit__()

    def test_first_ask_reuses_the_old_single_score(self):
        """Going from 1 ask to 3 pays only for the 2 new asks per company."""
        install_fake_openai()
        p = TempProject(config_overrides={"openai": {"alignment_votes": 1}})
        p.__enter__()
        try:
            p.seed_geocodes()
            p.run_enrich()
            FakeOpenAI.reset()
            self.assertEqual(p.run_match(["--faculty", "Ioulia Valla"]), 0)
            one = FakeOpenAI.counts().get("alignment", 0)
            self.assertGreater(one, 0)
            cfg_path = p.dir / "config.yaml"
            import yaml
            cfg = yaml.safe_load(cfg_path.read_text())
            cfg["openai"]["alignment_votes"] = 3
            cfg_path.write_text(yaml.safe_dump(cfg))
            FakeOpenAI.reset()
            self.assertEqual(p.run_match(["--faculty", "Ioulia Valla"]), 0)
            broken = unreadable(p.outputs()[0])
            self.assertEqual(FakeOpenAI.counts().get("alignment", 0), 2 * one + broken)
        finally:
            p.__exit__()


if __name__ == "__main__":
    unittest.main()
