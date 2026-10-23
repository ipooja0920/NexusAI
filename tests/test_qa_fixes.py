"""Fixes from the QA of Julie Burland's third run:
  1 alignment: no named link scores 1-2 (not 3); Mid / Large floor of 4
  2 animal-health makers are companies, not clinics
  3 companies with no description get no pre-filter seat
  4 Capital IQ code-like stub descriptions count as no description
  5 clinical-trial "university" partners without companies (CROs, drug makers)
  6 a parent and its subsidiaries are listed once
"""
import tempfile
import unittest
from pathlib import Path

import openpyxl
import pandas as pd

from match_faculty import (company_text, dedupe_across_tiers, prepare_companies, run_pass,
                           track_record_text)
from nexus.cache import Caches
from nexus.groups import build_groups, group_key, pick_listed
from nexus.organisation import CLINIC, COMPANY, academic_partners, organisation_type
from nexus.prefilter import company_profile_text, is_stub_description
from nexus.scoring import ALIGNMENT_SYSTEM_PROMPT
from nexus.settings import Settings
from tests.harness import TempProject, install_fake_openai


class TestAlignmentPrompt(unittest.TestCase):

    def test_no_link_is_one_or_two(self):
        self.assertIn("the score is 1-2, not 3", ALIGNMENT_SYSTEM_PROMPT)
        self.assertIn("names something specific", ALIGNMENT_SYSTEM_PROMPT)

    def test_floor_of_four(self):
        from match_faculty import tier_min_alignment
        cfg = Settings()
        self.assertEqual([tier_min_alignment(cfg, t) for t in ("Small", "Mid", "Large")],
                         [0, 4, 4])


class TestVeterinary(unittest.TestCase):

    def test_animal_health_maker_is_a_company(self):
        self.assertEqual(organisation_type(
            "Phibro Animal Health Corporation (NasdaqGM:PAHC)",
            "Phibro Animal Health Corporation operates as an animal health and mineral "
            "nutrition company in the United States."), COMPANY)
        self.assertEqual(organisation_type(
            "Animol Discovery, Inc.",
            "Animol Discovery, Inc. operates as a biopharmaceutical company focused on "
            "discovering novel medicines for animals."), COMPANY)

    def test_vet_care_settings_are_clinics(self):
        self.assertEqual(organisation_type("Absecon Veterinary Hospital", "-"), CLINIC)
        self.assertEqual(organisation_type("Animal & Bird Health Care Center & Hospital", "-"),
                         CLINIC)
        self.assertEqual(organisation_type(
            "Health Pet", "Health Pet operates veterinary medicine facilities."), CLINIC)
        self.assertEqual(organisation_type(
            "Elizabeth C Shurland Dvm", "Provides As An Animal Specialties Veterinarian."), CLINIC)


class TestCareProvidersFromRun4(unittest.TestCase):
    """Providers that were listed as companies in the Mid tab."""

    def test_doing_business_as_medical_center(self):
        from nexus.organisation import HOSPITAL
        self.assertEqual(organisation_type(
            "Southern Maine Health Care",
            "Southern Maine Health Care, doing business as Southern Maine Medical Center, "
            "provides medical care and services."), HOSPITAL)

    def test_psychotherapy_services(self):
        self.assertEqual(organisation_type(
            "The Center For Behavioral Health Inc.",
            "The Center for Behavioral Health, Inc. provides individual, family, and group "
            "psychotherapy services."), CLINIC)
        self.assertEqual(organisation_type(
            "Talkspace, Inc.",
            "Talkspace, Inc., together with its subsidiaries, operates as a virtual behavioral "
            "healthcare company that connects patients with licensed mental health providers "
            "in the United States. The company provides psychotherapy and psychiatry services "
            "to treat various conditions."), CLINIC)

    def test_software_company_stays_company(self):
        self.assertEqual(organisation_type(
            "CareCloud, Inc.", "CareCloud, Inc., a healthcare information technology company, "
            "provides technology-enabled business solutions, Software-as-a-Service offerings, "
            "and related business services to healthcare providers."), COMPANY)


class TestStubDescriptions(unittest.TestCase):

    def test_code_like_stubs(self):
        for t in ("Ofcs Clns Of Mdl D", "Sklld Nrsng Cr Fclt", "Drgs Prprtrs Sndr",
                  "Hlth Alld Svcs Ne", "-", "", None, float("nan")):
            self.assertTrue(is_stub_description(t), t)

    def test_real_short_descriptions_kept(self):
        for t in ("Provides Home Health Care Services.", "Operates A Kidney Dialysis Center.",
                  "Medical Hosp Equip", "General Hospital",
                  "Alpharma US Inc. operates as a subsidiary of Pfizer Inc.",
                  "ZetrOZ develops wearable ultrasound devices for musculoskeletal healing."):
            self.assertFalse(is_stub_description(t), t)

    def test_stub_is_no_text(self):
        self.assertEqual(company_text({"company_profile": "Ofcs Clns Of Mdl D",
                                       "description": "Ofcs Clns Of Mdl D"}), "")
        self.assertEqual(company_profile_text("Sklld Nrsng Cr Fclt", None), "")
        self.assertEqual(company_text({"company_profile": "Makes knee braces."}),
                         "Makes knee braces.")


class TestAcademicPartners(unittest.TestCase):

    def test_companies_taken_out(self):
        self.assertEqual(academic_partners(
            "Ironshore Pharmaceuticals and Development, Inc; Massachusetts General Hospital; "
            "Clinilabs, Inc.; American Cancer Society, Inc. (+3 more)"),
            "Massachusetts General Hospital; American Cancer Society, Inc. (+3 more)")
        self.assertIsNone(academic_partners("SCRI Development Innovations, LLC"))
        self.assertIsNone(academic_partners(None))
        self.assertEqual(academic_partners("UCLA; Duke University"), "UCLA; Duke University")

    def test_prepare_companies(self):
        c = [{"company": "Collegium", "description": "Collegium develops pain medicines. "
              "Collegium operates as a subsidiary of Big Pharma Corp.",
              "organisation_type": "Clinic / care provider", "parent_company": None,
              "ct_university_past": "Ironshore Pharmaceuticals, Inc",
              "ct_university_current": "Massachusetts General Hospital"}]
        prepare_companies(c)
        self.assertEqual(c[0]["organisation_type"], COMPANY)
        self.assertEqual(c[0]["parent_company"], "Big Pharma Corp")
        self.assertIsNone(c[0]["ct_university_past"])
        text = track_record_text(c[0])
        self.assertIn("Current academic and non-profit partners on clinical trials: "
                      "Massachusetts General Hospital", text)
        self.assertNotIn("Ironshore", text)


def _co(name, parent=None, dist=0.5):
    return {"company": name, "parent_company": parent, "distance_score": dist,
            "employee_score": 0.1, "revenue_score": 0.5, "organisation_type": "Company"}


class TestGroups(unittest.TestCase):

    def test_group_key(self):
        self.assertEqual(group_key("Lantheus Holdings, Inc. (NasdaqGM:LNTH)"), "lantheu")
        self.assertEqual(group_key("Integra LifeSciences Holdings Corporation (NasdaqGS:IART)"),
                         group_key("Integra LifeSciences Corporation"))
        self.assertEqual(group_key("Novo Nordisk A/S"), group_key("Novo Nordisk Inc."))

    def test_build_groups(self):
        cs = [_co("Lantheus Holdings, Inc. (NasdaqGM:LNTH)"),
              _co("Lantheus Medical Imaging, Inc.", parent="Lantheus Holdings, Inc"),
              _co("Integra LifeSciences Corporation"),
              _co("Integra LifeSciences Holdings Corporation"),
              _co("Sanofi-Aventis U.S. LLC", parent="Sanofi"),
              _co("Sanofi US Services Inc.", parent="Sanofi"),
              _co("Endodontics"), _co("Endodontics, L.L.C."),
              _co("Unrelated Devices Inc.")]
        g = build_groups(cs)
        self.assertEqual(g[0], g[1])
        self.assertEqual(g[2], g[3])
        self.assertEqual(g[4], g[5])
        self.assertNotEqual(g[0], g[2])
        self.assertIsNone(g[6])               # one-word names are not tied by name alone
        self.assertIsNone(g[8])

    def test_pick_listed(self):
        groups = [7, None, 7, None, 9]
        listed, gave = pick_listed([0, 1, 2, 3, 4], groups, 3)
        self.assertEqual(listed, [0, 1, 3])
        self.assertEqual(gave, {2: 0})
        listed, gave = pick_listed([0, 1, 2, 3, 4], groups, 5, blocked={9})
        self.assertEqual(listed, [0, 1, 3])
        self.assertEqual(gave[4], -1)

    def test_run_pass_lists_a_group_once(self):
        cfg = Settings()
        caches = Caches(Path(tempfile.mkdtemp()))
        cs = [_co("Parent Holdings Inc."), _co("Parent Sub LLC", parent="Parent Holdings Inc"),
              _co("Other Co")]
        align = {i: {"score": "8", "reason": ""} for i in range(3)}
        partner = {i: {"score": "7", "reason": ""} for i in range(3)}
        p = run_pass([0, 1, 2], cs, ["t"] * 3, align, partner, cfg, caches, client=None,
                     groups=build_groups(cs))
        self.assertEqual(len(p["final_companies"]), 2)
        self.assertIn(2, p["final_companies"])
        self.assertEqual(len(p["gave_way"]), 1)

    def test_dedupe_across_tiers_keeps_better_alignment(self):
        groups = [1, 1, None, None]
        overall = {0: 0.9, 1: 0.8, 2: 0.7, 3: 0.6}

        def pass_(order):
            return {"final_companies": list(order[:2]), "comp_order": list(order),
                    "final_providers": [], "prov_order": [], "gave_way": {},
                    "overall": overall, "final": list(order[:2])}
        widest = {"Mid": ("Mid", pass_([0, 2])), "Large": ("Large", pass_([1, 3]))}
        align = {0: {"score": "4"}, 1: {"score": "7"}, 2: {"score": "5"}, 3: {"score": "5"}}
        dedupe_across_tiers(widest, groups, align, 2)
        self.assertEqual(widest["Large"][1]["final_companies"], [1, 3])
        self.assertEqual(widest["Mid"][1]["final_companies"], [2])
        self.assertEqual(widest["Mid"][1]["gave_way"], {0: -1})


class TestPrefilterSeats(unittest.TestCase):

    def test_no_description_no_seat(self):
        install_fake_openai()
        p = TempProject(config_overrides={"prefilter": {"top_n_candidates": 3}})
        p.__enter__()
        try:
            p.seed_geocodes()
            p.run_enrich()
            cfg = p.settings()
            df = pd.read_excel(cfg.companies_enriched)
            small = df[df.revenue_usdmm.isna() | (df.revenue_usdmm < 100)].company.tolist()
            blank = small[0]
            for col in ("description", "company_profile", "research_focus"):
                if col in df:
                    df[col] = df[col].astype(object)
                    df.loc[df.company == blank, col] = "Ofcs Clns Of Mdl D"
            with pd.ExcelWriter(cfg.companies_enriched) as w:
                df.to_excel(w, index=False, sheet_name="Companies")
            self.assertEqual(p.run_match(), 0)
            ws = openpyxl.load_workbook(p.outputs()[0])["All Companies"]
            hdr = [c.value for c in ws[1]]
            rows = {ws.cell(r, 3).value: dict(zip(hdr, [c.value for c in ws[r]]))
                    for r in range(2, ws.max_row + 1)}
            self.assertEqual(rows[blank]["In Pre-filter"], "No")
            self.assertEqual(rows[blank]["Stage"], "Not scored (no description)")
        finally:
            p.__exit__()


if __name__ == "__main__":
    unittest.main()


class TestVerdictFirstOrder(unittest.TestCase):
    """A weak research match never outranks a strong one because it is near
    Storrs or large: tabs go Strong, then Moderate, then Weak, each by
    overall score (Axsome, alignment 4, was #11 in Leighton Core's Mid tab
    above Strong fits that are far away)."""

    def _co(self, name, dist, size):
        return {"company": name, "distance_score": dist, "employee_score": size,
                "revenue_score": size, "organisation_type": "Company"}

    def test_strong_before_weak_whatever_the_overall(self):
        from match_faculty import fit_band
        self.assertEqual([fit_band(s) for s in ("9", "7", "6", "5", "4", "1", "NA")],
                         [0, 0, 1, 1, 2, 2, 3])
        cfg = Settings()
        cfg.raw["funnel"]["final_top_n"] = 2
        caches = Caches(Path(tempfile.mkdtemp()))
        # the weak one is next door and large, the strong one far away and small
        cs = [self._co("Near Weak", 1.0, 1.0), self._co("Far Strong", 0.0, 0.0),
              self._co("Mid Moderate", 0.3, 0.5)]
        align = {0: {"score": "4", "reason": ""}, 1: {"score": "8", "reason": ""},
                 2: {"score": "5", "reason": ""}}
        partner = {i: {"score": "9", "reason": ""} for i in range(3)}
        p = run_pass([0, 1, 2], cs, ["t"] * 3, align, partner, cfg, caches, client=None)
        self.assertGreater(p["overall"][0], p["overall"][1])     # weak has the higher overall
        self.assertEqual(p["final_companies"], [1, 2])             # but strong, moderate are listed
        self.assertEqual(p["comp_order"], [1, 2, 0])
